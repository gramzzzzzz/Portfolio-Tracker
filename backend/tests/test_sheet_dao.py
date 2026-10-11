import datetime as dt
import io

import openpyxl
import pytest

from app.dao.base import LayoutError
from app.dao.xlsx.invt_sheet import InvtSheetDAO
from app.dao.xlsx.workbook import Workbook
from app.domain.models import NewLot

D = dt.date(2026, 10, 1)


def lots_all():
    return [
        NewLot("ITC LTD", "FMCG", D, 10, 260, 3.01, code="ITC", market_rate=260),
        NewLot("TATA MOTORS PV", "Automobile", D, 10, 276.5, 3.2, market_rate=276.5),
        NewLot("SBI ETF GOLD", "Gold ETF", D, 5, 110, 0.5, market_rate=101.68),
        NewLot("BHARTI AIRTEL", "Telecom", D, 2, 1900, 2.0, market_rate=1900),  # last block of section
        NewLot("NTPC", "Energy", D, 20, 350, 4.0, market_rate=350),  # brand new sector
        NewLot("PPFAS FLEXI", "MF", D, 1.234, 80, 0, market_rate=80),  # MF block types Value by hand
    ]


def test_read_sample(sheet_bytes):
    snap = InvtSheetDAO().read(sheet_bytes)
    assert snap.valuation_date == dt.date(2025, 10, 28)
    assert [l.name for l in snap.lots] == ["SBI ETF GOLD"]
    lot = snap.lots[0]
    assert (lot.qty, lot.rate, lot.charges, lot.cost, lot.sector, lot.block) == (60, 103, 1.24, 6181.24, "Gold ETF", "ETF")
    assert [s.name for s in snap.sectors][:3] == ["Automobile", "Bank", "FMCG"]
    assert snap.sector_to_block["Bank"] == "Banks" and snap.sector_to_block["MF"] == "Mutual Fund"
    assert snap.totals["grand_total_cost"] == pytest.approx(6181.24)
    assert len(snap.sections) == 3


def test_append_all_cases_and_verification(sheet_bytes):
    dao = InvtSheetDAO()
    out, res = dao.append(sheet_bytes, lots_all())
    assert all(c["passed"] for c in res.checks), res.checks
    names = {c["name"] for c in res.checks}
    assert {"existing_cells_intact", "no_overwrite", "grand_total_cost_delta", "summary_total_delta",
            "sector_total_delta", "new_rows_evaluate"} <= names
    snap = dao.read(out)
    assert len(snap.lots) == 7
    assert {s.name for s in snap.sectors} >= {"Energy"}
    assert "Energy" in [b.header for b in snap.blocks()]
    # totals reconcile exactly as in the original sheet (C60 / I57 style checks are 0)
    assert snap.totals["summary_total"] == pytest.approx(snap.totals["grand_total_cost"])
    assert snap.totals["sector_total"] == pytest.approx(snap.totals["grand_total_cost"])
    assert not [i for i in snap.issues if i.code in ("summary_mismatch", "sector_mismatch", "lot_not_in_summary")]


def test_new_rows_use_section_formulas(sheet_bytes):
    out, res = InvtSheetDAO().append(sheet_bytes, lots_all())
    cells = Workbook(out).sheet("Shares").cells()
    lot_rows = {r.ref or r.description: r.row for r in res.rows if r.kind == "lot"}
    itc = next(r.row for r in res.rows if r.kind == "lot" and "ITC" in r.description)
    assert cells[(itc, 7)].formula == f"E{itc}*F{itc}"  # G = qty*rate
    assert cells[(itc, 9)].formula == f"G{itc}+H{itc}"  # I = value + comn
    assert cells[(itc, 11)].formula == f"E{itc}"  # K = qty
    assert cells[(itc, 19)].formula == f"($S$1-C{itc})/365"
    assert cells[(itc, 8)].value == 3.01 and cells[(itc, 8)].formula is None  # H typed value
    gold = [r.row for r in res.rows if r.kind == "lot" and "GOLD" in r.description][0]
    assert cells[(gold, 18)].formula == f"P{gold}/G{gold}"  # ETF section pattern differs (P/G)
    mf = [r.row for r in res.rows if r.kind == "lot" and "PPFAS" in r.description][0]
    assert cells[(mf, 7)].formula is None and cells[(mf, 7)].value == pytest.approx(98.72)
    assert lot_rows


def test_summary_block_shifted_not_overwritten(sheet_bytes):
    before = Workbook(sheet_bytes).sheet("Shares").cells()
    out, res = InvtSheetDAO().append(sheet_bytes, lots_all()[:1])
    after = Workbook(out).sheet("Shares").cells()
    # the 'Summary' label moved down by the number of rows inserted above it
    old = next(r for (r, c), v in before.items() if v.value == "Summary")
    new = next(r for (r, c), v in after.items() if v.value == "Summary")
    assert new == old + 1
    assert after[(new + 4, 1)].value == "SBI ETF GOLD"  # still in its own summary row


def test_existing_stock_goes_to_its_lot_group(sheet_bytes):
    dao = InvtSheetDAO()
    out, res = dao.append(sheet_bytes, [NewLot("sbi etf gold", "Gold ETF", D, 1, 100, 0.1)])
    snap = dao.read(out)
    rows = [l.row for l in snap.lots if l.name.lower() == "sbi etf gold"]
    assert rows == [31, 32]
    assert [r.kind for r in res.rows] == ["lot"]  # no new summary row: it already exists


def test_workbook_parts_preserved(sheet_bytes):
    import zipfile
    out, _ = InvtSheetDAO().append(sheet_bytes, lots_all()[:2])
    zi, zo = zipfile.ZipFile(io.BytesIO(sheet_bytes)), zipfile.ZipFile(io.BytesIO(out))
    untouched = ["xl/styles.xml", "xl/theme/theme1.xml", "xl/printerSettings/printerSettings1.bin", "docProps/core.xml"]
    for n in untouched:
        assert zi.read(n) == zo.read(n)
    assert "xl/calcChain.xml" not in zo.namelist()
    wb = openpyxl.load_workbook(io.BytesIO(out))
    assert wb.calculation.fullCalcOnLoad
    ws = wb["Shares"]
    assert ws.auto_filter.ref == "A1:P41"  # grew with the two inserted rows
    assert {str(m) for m in ws.merged_cells.ranges} >= {"E1:I1", "K1:O1"}


def test_dry_run_returns_no_bytes(sheet_bytes):
    out, res = InvtSheetDAO().append(sheet_bytes, lots_all()[:1], dry_run=True)
    assert out is None and res.rows


def test_rejects_unrelated_workbook():
    wb = openpyxl.Workbook()
    wb.active["A1"] = "hello"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(LayoutError):
        InvtSheetDAO().read(buf.getvalue())


def test_sequential_appends_stay_consistent(sheet_bytes):
    dao = InvtSheetDAO()
    data = sheet_bytes
    for i, nl in enumerate(lots_all()):
        data, res = dao.append(data, [nl])
        assert all(c["passed"] for c in res.checks), (i, res.checks)
    snap = dao.read(data)
    assert snap.totals["summary_total"] == pytest.approx(snap.totals["grand_total_cost"])


def test_new_sector_added_everywhere_once(sheet_bytes):
    """A new sector gets a heading, a Summary row per stock and exactly one Sector-table row,
    and later stocks of that sector reuse them (in the same batch and in later imports)."""
    dao = InvtSheetDAO()
    data, r1 = dao.append(sheet_bytes, [NewLot("NTPC", "Energy", D, 20, 350, 4.0)])
    assert [p.kind for p in r1.rows] == ["block_header", "lot", "summary+sector"]
    data, r2 = dao.append(data, [NewLot("POWER GRID", "Energy", D, 10, 300, 2.0)])
    assert [p.kind for p in r2.rows] == ["lot", "summary"]  # no second heading / sector row
    data, r3 = dao.append(data, [NewLot("DLF", "Realty", D, 5, 800, 2.0), NewLot("GODREJ PROP", "Realty", D, 2, 2500, 2.0),
                                 NewLot("SILVERBEES", "Silver ETF", D, 10, 95, 0.5)])
    for r in (r1, r2, r3):
        assert all(c["passed"] for c in r.checks), r.checks
    snap = dao.read(data)
    headings = [b.header for b in snap.blocks()]
    assert headings.count("Energy") == 1 and headings.count("Realty") == 1
    assert "Silver ETF" not in headings  # an ETF-like sector is placed under the existing ETF heading
    assert {l.name: l.block for l in snap.lots}["SILVERBEES"] == "ETF"
    names = [s.name for s in snap.sectors]
    assert names.count("Energy") == 1 and names.count("Realty") == 1 and names.count("Silver ETF") == 1
    by_sector = {s.name: s.amount for s in snap.sectors}
    assert by_sector["Energy"] == pytest.approx(7004 + 3002) and by_sector["Realty"] == pytest.approx(4002 + 5002)
    assert sum(s.pct for s in snap.sectors) == pytest.approx(1.0)  # %Age column includes the new rows
    assert snap.totals["sector_total"] == pytest.approx(snap.totals["grand_total_cost"])
    assert snap.totals["summary_total"] == pytest.approx(snap.totals["grand_total_cost"])
    assert not [i for i in snap.issues if i.severity != "info"]
