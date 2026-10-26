"""Surgical OOXML (xlsx) workbook editor.

Why not openpyxl for writing?  openpyxl re-serialises the whole workbook, silently drops
parts it does not understand and does NOT rewrite formulas when rows are inserted. This
module edits the package XML in place: only the parts that must change are touched and
every other byte (styles, printer settings, theme, custom props...) is preserved.

Supported operation: inserting whole rows into one worksheet, exactly as Excel's
"Insert row" does (cells, formulas in every sheet, defined names, merged cells,
conditional formatting, data validation, autofilter, tables, hyperlinks, row breaks and
the sheet dimension are all shifted), and writing values/formulas into the new cells.
"""
from __future__ import annotations

import copy
import datetime as dt
import io
import posixpath
import re
import zipfile
from dataclasses import dataclass
from typing import Any, Iterable

from lxml import etree
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from openpyxl.utils import column_index_from_string, get_column_letter

from . import refs
from .evaluator import BLANK, Evaluator, Unsupported, XLError

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_RNS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
N = {"m": NS}


def q(tag: str) -> str:
    return f"{{{NS}}}{tag}"


EXCEL_EPOCH = dt.datetime(1899, 12, 30)


def to_serial(d: dt.date | dt.datetime) -> float:
    if not isinstance(d, dt.datetime):
        d = dt.datetime(d.year, d.month, d.day)
    delta = d - EXCEL_EPOCH
    return delta.days + delta.seconds / 86400


def from_serial(v: float) -> dt.datetime:
    return EXCEL_EPOCH + dt.timedelta(days=float(v))


@dataclass
class Cell:
    row: int
    col: int
    value: Any  # float | str | bool | XLError | BLANK
    formula: str | None
    style: int
    is_date: bool

    @property
    def coord(self) -> str:
        return refs.cell_name(self.row, self.col)


class Workbook:
    """Holds an xlsx package in memory; edits a single target worksheet."""

    def __init__(self, data: bytes):
        self._zip_order: list[str] = []
        self._parts: dict[str, bytes] = {}
        self._zinfo: dict[str, zipfile.ZipInfo] = {}
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for info in z.infolist():
                self._zip_order.append(info.filename)
                self._parts[info.filename] = z.read(info.filename)
                self._zinfo[info.filename] = info
        self._trees: dict[str, etree._ElementTree] = {}
        self._dirty: set[str] = set()
        self._removed: set[str] = set()
        self.sheet_paths = self._resolve_sheets()
        self._load_shared_strings()
        self._load_styles()

    # ------------------------------------------------------------------ package
    def _tree(self, name: str) -> etree._ElementTree:
        if name not in self._trees:
            parser = etree.XMLParser(remove_blank_text=False, resolve_entities=False, huge_tree=True)
            self._trees[name] = etree.ElementTree(etree.fromstring(self._parts[name], parser))
        return self._trees[name]

    def _touch(self, name: str):
        self._dirty.add(name)

    @staticmethod
    def _rels_path(part: str) -> str:
        d, f = posixpath.split(part)
        return posixpath.join(d, "_rels", f + ".rels")

    def _rels(self, part: str) -> dict[str, tuple[str, str]]:
        """rId -> (type, absolute target)"""
        rp = self._rels_path(part)
        if rp not in self._parts:
            return {}
        root = self._tree(rp).getroot()
        out = {}
        base = posixpath.dirname(part)
        for r in root:
            tgt = r.get("Target")
            if r.get("TargetMode") == "External":
                continue
            abs_t = tgt.lstrip("/") if tgt.startswith("/") else posixpath.normpath(posixpath.join(base, tgt))
            out[r.get("Id")] = (r.get("Type"), abs_t)
        return out

    def _resolve_sheets(self) -> dict[str, str]:
        wb = self._tree("xl/workbook.xml").getroot()
        rels = self._rels("xl/workbook.xml")
        out = {}
        for s in wb.find("m:sheets", N):
            rid = s.get(f"{{{RNS}}}id")
            out[s.get("name")] = rels[rid][1]
        return out

    def mark_full_calc(self) -> None:
        """Ask Excel / LibreOffice to recalculate every formula when the file is opened."""
        wbroot = self._tree("xl/workbook.xml").getroot()
        calc = wbroot.find("m:calcPr", N)
        if calc is None:
            calc = etree.SubElement(wbroot, q("calcPr"))
        calc.set("fullCalcOnLoad", "1")
        self._touch("xl/workbook.xml")

    def add_part(self, name: str, data: bytes) -> None:
        """Register a new package part (written at the end of the zip)."""
        if name not in self._parts:
            self._zip_order.append(name)
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.external_attr = 0o600 << 16
            self._zinfo[name] = info
        self._parts[name] = data

    def sheet_names(self) -> list[str]:
        return list(self.sheet_paths)

    def to_bytes(self) -> bytes:
        for name in self._dirty:
            if name in self._trees:
                self._parts[name] = etree.tostring(
                    self._trees[name].getroot(), xml_declaration=True, encoding="UTF-8", standalone=True
                )
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for name in self._zip_order:
                if name in self._removed:
                    continue
                info = zipfile.ZipInfo(name, date_time=self._zinfo[name].date_time)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = self._zinfo[name].external_attr
                z.writestr(info, self._parts[name])
        return buf.getvalue()

    # ------------------------------------------------------------- shared strings
    def _load_shared_strings(self):
        rels = self._rels("xl/workbook.xml")
        self._sst_path = None
        for typ, tgt in rels.values():
            if typ.endswith("/sharedStrings"):
                self._sst_path = tgt
        self._strings: list[str] = []
        if self._sst_path and self._sst_path in self._parts:
            for si in self._tree(self._sst_path).getroot().findall("m:si", N):
                self._strings.append("".join(si.itertext()))

    def _add_shared_string(self, text: str) -> int:
        if self._sst_path is None:
            raise RuntimeError("Workbook has no shared strings part")  # pragma: no cover
        root = self._tree(self._sst_path).getroot()
        if text in self._strings:
            idx = self._strings.index(text)
        else:
            si = etree.SubElement(root, q("si"))
            t = etree.SubElement(si, q("t"))
            t.text = text
            if text != text.strip():
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            self._strings.append(text)
            idx = len(self._strings) - 1
            root.set("uniqueCount", str(len(self._strings)))
        if root.get("count") is not None:
            root.set("count", str(int(root.get("count")) + 1))
        self._touch(self._sst_path)
        return idx

    # --------------------------------------------------------------------- styles
    def _load_styles(self):
        self._date_styles: set[int] = set()
        path = "xl/styles.xml"
        if path not in self._parts:
            return
        root = self._tree(path).getroot()
        custom = {}
        nf = root.find("m:numFmts", N)
        if nf is not None:
            for f in nf:
                custom[int(f.get("numFmtId"))] = f.get("formatCode")
        xfs = root.find("m:cellXfs", N)
        if xfs is None:
            return
        for i, xf in enumerate(xfs):
            fid = int(xf.get("numFmtId", "0"))
            code = custom.get(fid) or BUILTIN_FORMATS.get(fid)
            if code and is_date_format(code):
                self._date_styles.add(i)

    def is_date_style(self, style: int) -> bool:
        return style in self._date_styles

    # ---------------------------------------------------------------- sheet model
    def sheet(self, name: str) -> "Sheet":
        return Sheet(self, name, self.sheet_paths[name])


class Sheet:
    def __init__(self, wb: Workbook, name: str, path: str):
        self.wb = wb
        self.name = name
        self.path = path
        self.root = wb._tree(path).getroot()
        self.sheet_data = self.root.find("m:sheetData", N)
        self._unshare_formulas()

    # ------------------------------------------------------------------ reading
    def rows(self) -> Iterable[etree._Element]:
        return self.sheet_data.findall("m:row", N)

    def _row_el(self, r: int):
        for row in self.sheet_data.findall("m:row", N):
            rr = int(row.get("r"))
            if rr == r:
                return row
            if rr > r:
                return None
        return None

    def max_row(self) -> int:
        rows = self.sheet_data.findall("m:row", N)
        return int(rows[-1].get("r")) if rows else 0

    def max_col(self) -> int:
        mc = 1
        for c in self.sheet_data.iter(q("c")):
            mc = max(mc, refs.split_cell(c.get("r"))[1])
        return mc

    def _cell_value(self, c: etree._Element):
        t = c.get("t", "n")
        v = c.find("m:v", N)
        if t == "inlineStr":
            is_ = c.find("m:is", N)
            return "".join(is_.itertext()) if is_ is not None else BLANK
        if v is None or v.text is None:
            return BLANK
        txt = v.text
        if t == "s":
            return self.wb._strings[int(txt)]
        if t == "str":
            return txt
        if t == "b":
            return txt == "1"
        if t == "e":
            return XLError(txt)
        try:
            return float(txt)
        except ValueError:
            return txt

    def cells(self) -> dict[tuple[int, int], Cell]:
        out = {}
        for row in self.rows():
            for c in row.findall("m:c", N):
                r, col = refs.split_cell(c.get("r"))
                f = c.find("m:f", N)
                style = int(c.get("s", "0"))
                out[(r, col)] = Cell(
                    row=r,
                    col=col,
                    value=self._cell_value(c),
                    formula=(f.text or "") if f is not None else None,
                    style=style,
                    is_date=self.wb.is_date_style(style),
                )
        return out

    def merged_ranges(self) -> list[str]:
        return [m.get("ref") for m in self.root.iterfind("m:mergeCells/m:mergeCell", N)]

    def row_attrs(self, r: int) -> dict[str, str]:
        el = self._row_el(r)
        return dict(el.attrib) if el is not None else {}

    # ---------------------------------------------------------- shared formulas
    def _unshare_formulas(self):
        """Expand shared formulas to explicit per-cell formulas (semantics unchanged).
        Needed because row insertion can split shared ranges."""
        masters: dict[str, tuple[str, str]] = {}
        for c in self.sheet_data.iter(q("c")):
            f = c.find("m:f", N)
            if f is not None and f.get("t") == "shared" and f.get("ref") is not None and f.text:
                masters[f.get("si")] = (f.text, c.get("r"))
        if not masters:
            return
        for c in self.sheet_data.iter(q("c")):
            f = c.find("m:f", N)
            if f is None or f.get("t") != "shared":
                continue
            si = f.get("si")
            text, origin = masters[si]
            f.text = text if origin == c.get("r") else refs.translate(text, origin, c.get("r"))
            for a in ("t", "ref", "si"):
                if a in f.attrib:
                    del f.attrib[a]
        self.wb._touch(self.path)

    # ------------------------------------------------------------ evaluation
    def evaluator(self) -> Evaluator:
        cells = self.cells()

        def get(r, c):
            cell = cells.get((r, c))
            if cell is None:
                return BLANK, None
            return cell.value, cell.formula

        return Evaluator(get, self.max_row(), self.max_col())

    def refresh_cached_values(self) -> list[str]:
        """Recompute every formula with the built-in evaluator and store the result as the
        cached value. Cells the evaluator cannot handle get their cached value removed
        (Excel recalculates everything on open thanks to fullCalcOnLoad).
        Returns coordinates that could not be evaluated."""
        ev = self.evaluator()
        failed = []
        for c in self.sheet_data.iter(q("c")):
            f = c.find("m:f", N)
            if f is None:
                continue
            r, col = refs.split_cell(c.get("r"))
            v = c.find("m:v", N)
            try:
                res = ev.value(r, col)
            except (Unsupported, RecursionError, ValueError, TypeError):
                failed.append(c.get("r"))
                if v is not None:
                    c.remove(v)
                if c.get("t") in ("e", "str", "b"):
                    del c.attrib["t"]
                continue
            if v is None:
                v = etree.SubElement(c, q("v"))
            if isinstance(res, XLError):
                c.set("t", "e")
                v.text = str(res)
            elif isinstance(res, bool):
                c.set("t", "b")
                v.text = "1" if res else "0"
            elif isinstance(res, str):
                c.set("t", "str")
                v.text = res
            else:
                if "t" in c.attrib:
                    del c.attrib["t"]
                v.text = repr(float(res)) if not float(res).is_integer() else str(int(res))
        self.wb._touch(self.path)
        return failed

    # ---------------------------------------------------------------- insertion
    def insert_rows(self, spec: refs.InsertSpec):
        """Insert ``spec.count`` empty rows at ``spec.at`` shifting everything below."""
        wb = self.wb
        # 1. shift rows & cells of this sheet (iterate bottom-up to keep order valid)
        for row in reversed(self.sheet_data.findall("m:row", N)):
            r = int(row.get("r"))
            if r < spec.at:
                break
            row.set("r", str(r + spec.count))
            for c in row.findall("m:c", N):
                rr, cc = refs.split_cell(c.get("r"))
                c.set("r", refs.cell_name(rr + spec.count, cc))
        # 2. formulas in every sheet (shared formulas elsewhere are expanded first so that
        #    relative children stay correct)
        for sname, spath in wb.sheet_paths.items():
            if spath != self.path:
                Sheet(wb, sname, spath)  # constructor unshares formulas in place
            root = wb._tree(spath).getroot() if spath != self.path else self.root
            touched = False
            for f in root.iter(q("f")):
                if f.text:
                    frow = -1
                    if spath == self.path:
                        cur, _ = refs.split_cell(f.getparent().get("r"))
                        frow = cur - spec.count if cur >= spec.at + spec.count else cur
                    nt = refs.shift_formula(
                        f.text, spec, own_sheet=sname, target_sheet=self.name, formula_row=frow
                    )
                    if nt != f.text:
                        f.text = nt
                        touched = True
                if f.get("ref") and spath == self.path:  # array formula range
                    f.set("ref", refs.shift_sqref(f.get("ref"), spec))
                    touched = True
            # conditional formatting / data validation formulas
            for tag in ("formula", "formula1", "formula2"):
                for f in root.iter(q(tag)):
                    if f.text:
                        nt = refs.shift_formula(f.text, spec, own_sheet=sname, target_sheet=self.name)
                        if nt != f.text:
                            f.text = nt
                            touched = True
            if touched:
                wb._touch(spath)
        # 3. sheet-level ranges
        root = self.root
        dim = root.find("m:dimension", N)
        if dim is not None:
            dim.set("ref", self._grow_dimension(dim.get("ref"), spec))
        for mc in root.iterfind("m:mergeCells/m:mergeCell", N):
            mc.set("ref", refs.shift_sqref(mc.get("ref"), spec))
        for cf in root.iterfind("m:conditionalFormatting", N):
            cf.set("sqref", refs.shift_sqref(cf.get("sqref"), spec, spec.extend_below))
        for dv in root.iterfind("m:dataValidations/m:dataValidation", N):
            dv.set("sqref", refs.shift_sqref(dv.get("sqref"), spec, spec.extend_below))
        af = root.find("m:autoFilter", N)
        if af is not None:
            af.set("ref", refs.shift_sqref(af.get("ref"), spec, spec.extend_below))
        for hl in root.iterfind("m:hyperlinks/m:hyperlink", N):
            hl.set("ref", refs.shift_sqref(hl.get("ref"), spec))
        for brk in root.iterfind("m:rowBreaks/m:brk", N):
            b = int(brk.get("id"))
            if b >= spec.at:
                brk.set("id", str(b + spec.count))
        # 4. tables attached to this sheet
        for typ, tgt in wb._rels(self.path).values():
            if typ.endswith("/table") and tgt in wb._parts:
                troot = wb._tree(tgt).getroot()
                troot.set("ref", refs.shift_sqref(troot.get("ref"), spec))
                taf = troot.find("m:autoFilter", N)
                if taf is not None:
                    taf.set("ref", refs.shift_sqref(taf.get("ref"), spec))
                wb._touch(tgt)
        # 5. workbook defined names
        wbroot = wb._tree("xl/workbook.xml").getroot()
        for dn in wbroot.iterfind("m:definedNames/m:definedName", N):
            if dn.text:
                own = None
                if dn.get("localSheetId") is not None:
                    own = wb.sheet_names()[int(dn.get("localSheetId"))]
                nt = refs.shift_formula(dn.text, spec, own_sheet=own or "\0", target_sheet=self.name)
                if nt != dn.text:
                    dn.text = nt
        wb.mark_full_calc()
        self._drop_calc_chain()
        self.wb._touch(self.path)

    @staticmethod
    def _grow_dimension(ref: str, spec: refs.InsertSpec) -> str:
        if ":" not in ref:
            return ref
        a, b = ref.split(":")
        ra, ca = refs.split_cell(a)
        rb, cb = refs.split_cell(b)
        if rb >= spec.at:
            rb += spec.count
        return f"{refs.cell_name(ra, ca)}:{refs.cell_name(rb, cb)}"

    def _drop_calc_chain(self):
        wb = self.wb
        path = "xl/calcChain.xml"
        if path not in wb._parts or path in wb._removed:
            return
        wb._removed.add(path)
        rels_root = wb._tree("xl/_rels/workbook.xml.rels").getroot()
        for r in list(rels_root):
            if r.get("Type", "").endswith("/calcChain"):
                rels_root.remove(r)
        wb._touch("xl/_rels/workbook.xml.rels")
        ct = wb._tree("[Content_Types].xml").getroot()
        for o in list(ct):
            if o.get("PartName") == "/xl/calcChain.xml":
                ct.remove(o)
        wb._touch("[Content_Types].xml")

    # ------------------------------------------------------------------ writing
    def create_row(self, r: int, template_row: int | None) -> etree._Element:
        """Create (empty) row element r, copying row attributes and cell styles from template."""
        if self._row_el(r) is not None and len(self._row_el(r)):
            raise RuntimeError(f"Row {r} is not empty")
        existing = self._row_el(r)
        if existing is not None:
            self.sheet_data.remove(existing)
        new = etree.Element(q("row"))
        new.set("r", str(r))
        tmpl = self._row_el(template_row) if template_row else None
        if tmpl is not None:
            for k, v in tmpl.attrib.items():
                if k != "r":
                    new.set(k, v)
            for c in tmpl.findall("m:c", N):
                _, col = refs.split_cell(c.get("r"))
                nc = etree.SubElement(new, q("c"))
                nc.set("r", refs.cell_name(r, col))
                if c.get("s"):
                    nc.set("s", c.get("s"))
        # insert in order
        after = None
        for row in self.sheet_data.findall("m:row", N):
            if int(row.get("r")) < r:
                after = row
            else:
                break
        if after is None:
            self.sheet_data.insert(0, new)
        else:
            after.addnext(new)
        self.wb._touch(self.path)
        return new

    def _cell_el(self, r: int, col: int, create=True):
        row = self._row_el(r)
        if row is None:
            if not create:
                return None
            row = self.create_row(r, None)
        coord = refs.cell_name(r, col)
        prev = None
        for c in row.findall("m:c", N):
            _, cc = refs.split_cell(c.get("r"))
            if cc == col:
                return c
            if cc < col:
                prev = c
        if not create:
            return None
        nc = etree.Element(q("c"))
        nc.set("r", coord)
        if prev is None:
            row.insert(0, nc)
        else:
            prev.addnext(nc)
        return nc

    def set_cell(self, r: int, col: int, value: Any = None, *, formula: str | None = None, style: int | None = None):
        c = self._cell_el(r, col)
        for child in list(c):
            c.remove(child)
        if "t" in c.attrib:
            del c.attrib["t"]
        if style is not None:
            c.set("s", str(style))
        if formula is not None:
            f = etree.SubElement(c, q("f"))
            f.text = formula.lstrip("=")
        elif value is None or value is BLANK:
            pass
        elif isinstance(value, bool):
            c.set("t", "b")
            etree.SubElement(c, q("v")).text = "1" if value else "0"
        elif isinstance(value, (int, float)):
            etree.SubElement(c, q("v")).text = repr(float(value)) if not float(value).is_integer() else str(int(value))
        elif isinstance(value, (dt.date, dt.datetime)):
            serial = to_serial(value)
            etree.SubElement(c, q("v")).text = str(int(serial)) if float(serial).is_integer() else repr(serial)
        elif isinstance(value, str):
            idx = self.wb._add_shared_string(value)
            c.set("t", "s")
            etree.SubElement(c, q("v")).text = str(idx)
        else:
            raise TypeError(f"Unsupported value type {type(value)}")
        self.wb._touch(self.path)

    def cell_style(self, r: int, col: int) -> int | None:
        c = self._cell_el(r, col, create=False)
        return int(c.get("s")) if c is not None and c.get("s") else None


def col_index(letter: str) -> int:
    return column_index_from_string(letter)


def col_letter(idx: int) -> str:
    return get_column_letter(idx)
