"""App-owned worksheets inside the tracking workbook (e.g. ``Transactions``, ``Gold``).

The owner's rule: these sheets must not affect the original tracking sheet in any way.
This module therefore:

* never opens or re-serialises any other worksheet part (the original sheet's XML, the
  shared-strings table and the calc chain stay byte-identical - verified after every write);
* writes text as *inline strings* so ``sharedStrings.xml`` is not touched;
* only **appends** its own fonts / number formats / cell formats to ``styles.xml`` (existing
  style indexes, which the original sheet uses, keep their meaning);
* adds the new sheet to ``workbook.xml`` / its rels / ``[Content_Types].xml`` at the END of the
  sheet list, so existing sheet indexes (``localSheetId``) are unchanged;
* rewrites an app sheet as a whole from domain rows each time (the app owns its layout).

Reading accepts both inline and shared strings, so the owner may also edit these sheets in
Excel and upload the workbook; rows are parsed by header name, not position.
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass
from typing import Any, Iterable

from lxml import etree
from openpyxl.utils import get_column_letter

from .evaluator import BLANK, XLError
from .workbook import N, NS, RNS, Workbook, from_serial, q, to_serial

CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
WS_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
WS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
DATE_FMT = "dd-mmm-yyyy"
NUM_FMT = "#,##0.00"
QTY_FMT = "#,##0.####"

# column kinds
TEXT, DATE, NUM, QTY, INT = "text", "date", "num", "qty", "int"


class AppSheetError(Exception):
    pass


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    kind: str = TEXT
    width: float = 14.0
    required: bool = False


@dataclass(frozen=True)
class TableSpec:
    sheet: str
    columns: tuple[Column, ...]
    title: str  # line written in row 1 (explains the sheet); headers are in row 2

    @property
    def header_row(self) -> int:
        return 2


def _sig(el) -> tuple:
    """Namespace-independent structural signature of an XML element."""
    return (etree.QName(el).localname, tuple(sorted(el.attrib.items())), tuple(_sig(c) for c in el))


def _norm(s: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


# ============================================================================ reading
def read_table(wb: Workbook, spec: TableSpec) -> tuple[list[dict], list[str]]:
    """Rows of an app sheet as dicts keyed by column key. Missing sheet -> ([], []).
    Returns (rows, warnings). Never mutates the workbook."""
    if spec.sheet not in wb.sheet_paths:
        return [], []
    root = wb._tree(wb.sheet_paths[spec.sheet]).getroot()
    sd = root.find("m:sheetData", N)
    grid: dict[int, dict[int, Any]] = {}
    for row in sd.findall("m:row", N):
        r = int(row.get("r"))
        for c in row.findall("m:c", N):
            ref = c.get("r")
            col = 0
            for ch in re.match(r"[A-Z]+", ref).group(0):
                col = col * 26 + ord(ch) - 64
            grid.setdefault(r, {})[col] = _value(wb, c)
    # locate the header row (first row containing at least two known headers)
    want = {_norm(c.header): c for c in spec.columns}
    header_row, colmap = None, {}
    for r in sorted(grid)[:10]:
        m = {ci: want[_norm(v)] for ci, v in grid[r].items() if isinstance(v, str) and _norm(v) in want}
        if len(m) >= 2:
            header_row, colmap = r, m
            break
    warnings: list[str] = []
    if header_row is None:
        return [], [f"Sheet '{spec.sheet}' has no recognisable header row; it was ignored"]
    missing = [c.header for c in spec.columns if c.required and c not in colmap.values()]
    if missing:
        warnings.append(f"Sheet '{spec.sheet}' is missing column(s): {', '.join(missing)}")
    out = []
    for r in sorted(k for k in grid if k > header_row):
        cells = grid[r]
        rec: dict[str, Any] = {"_row": r}
        empty = True
        for ci, col in colmap.items():
            v = cells.get(ci, BLANK)
            v = _coerce(v, col.kind)
            if v is not None:
                empty = False
            rec[col.key] = v
        if not empty:
            out.append(rec)
    return out, warnings


def _value(wb: Workbook, c: etree._Element):
    t = c.get("t", "n")
    if t == "inlineStr":
        is_ = c.find("m:is", N)
        return "".join(is_.itertext()) if is_ is not None else BLANK
    v = c.find("m:v", N)
    if v is None or v.text is None:
        return BLANK
    if t == "s":
        return wb._strings[int(v.text)]
    if t in ("str", "e"):
        return v.text if t == "str" else XLError(v.text)
    if t == "b":
        return v.text == "1"
    try:
        return float(v.text)
    except ValueError:
        return v.text


def _coerce(v: Any, kind: str):
    if v is BLANK or v is None or isinstance(v, XLError):
        return None
    if kind == TEXT:
        if isinstance(v, float):
            return str(int(v)) if v.is_integer() else str(v)
        s = str(v).strip()
        return s or None
    if kind == DATE:
        if isinstance(v, (int, float)):
            return from_serial(float(v)).date()
        s = str(v).strip()
        for fmt in ("%d-%b-%Y", "%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%b-%y"):
            try:
                return dt.datetime.strptime(s, fmt).date()
            except ValueError:
                pass
        return None
    # numeric kinds
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(v) if kind == INT else float(v)
    try:
        f = float(str(v).replace(",", "").replace("₹", "").strip())
        return int(f) if kind == INT else f
    except ValueError:
        return None


# ============================================================================ writing
class _Styles:
    """Appends (or reuses) a few cell formats in styles.xml. Never edits existing entries."""

    def __init__(self, wb: Workbook):
        self.wb = wb
        self.root = wb._tree("xl/styles.xml").getroot()
        self._xf_cache: dict[tuple, int] = {}

    def _container(self, tag: str, before: tuple[str, ...]) -> etree._Element:
        el = self.root.find(f"m:{tag}", N)
        if el is None:
            el = etree.Element(q(tag))
            el.set("count", "0")
            anchor = None
            for b in before:
                anchor = self.root.find(f"m:{b}", N)
                if anchor is not None:
                    break
            if anchor is not None:
                anchor.addprevious(el)
            else:
                self.root.insert(0, el)
        return el

    def _append(self, parent: etree._Element, child: etree._Element) -> int:
        parent.append(child)
        n = len(parent)
        parent.set("count", str(n))
        return n - 1

    def numfmt(self, code: str) -> int:
        nf = self._container("numFmts", ("fonts",))
        ids = []
        for f in nf:
            if f.get("formatCode") == code:
                return int(f.get("numFmtId"))
            ids.append(int(f.get("numFmtId")))
        new_id = max(ids + [163]) + 1
        el = etree.Element(q("numFmt"))
        el.set("numFmtId", str(new_id))
        el.set("formatCode", code)
        self._append(nf, el)
        return new_id

    def font(self, bold: bool, color: str | None = None) -> int:
        fonts = self.root.find("m:fonts", N)
        el = etree.Element(q("font"))
        if bold:
            etree.SubElement(el, q("b"))
        etree.SubElement(el, q("sz")).set("val", "10")
        if color:
            etree.SubElement(el, q("color")).set("rgb", color)
        etree.SubElement(el, q("name")).set("val", "Arial")
        etree.SubElement(el, q("family")).set("val", "2")
        sig = _sig(el)
        for i, f in enumerate(fonts):
            if _sig(f) == sig:
                return i
        return self._append(fonts, el)

    def fill(self, rgb: str) -> int:
        fills = self.root.find("m:fills", N)
        el = etree.Element(q("fill"))
        pf = etree.SubElement(el, q("patternFill"))
        pf.set("patternType", "solid")
        etree.SubElement(pf, q("fgColor")).set("rgb", rgb)
        etree.SubElement(pf, q("bgColor")).set("indexed", "64")
        sig = _sig(el)
        for i, f in enumerate(fills):
            if _sig(f) == sig:
                return i
        return self._append(fills, el)

    def xf(self, *, num: int = 0, font: int = 0, fill: int = 0, wrap: bool = False) -> int:
        key = (num, font, fill, wrap)
        if key in self._xf_cache:
            return self._xf_cache[key]
        xfs = self.root.find("m:cellXfs", N)
        el = etree.Element(q("xf"))
        el.set("numFmtId", str(num))
        el.set("fontId", str(font))
        el.set("fillId", str(fill))
        el.set("borderId", "0")
        el.set("xfId", "0")
        if num:
            el.set("applyNumberFormat", "1")
        if font:
            el.set("applyFont", "1")
        if fill:
            el.set("applyFill", "1")
        if wrap:
            el.set("applyAlignment", "1")
            etree.SubElement(el, q("alignment")).set("wrapText", "1")
        sig = _sig(el)
        for i, x in enumerate(xfs):
            if _sig(x) == sig:
                self._xf_cache[key] = i
                return i
        idx = self._append(xfs, el)
        self._xf_cache[key] = idx
        return idx


def _escape_sheet_xml_text(s: str) -> str:
    # strip characters that are illegal in XML 1.0
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", s)


def _sheet_xml(spec: TableSpec, rows: list[dict], st: _Styles) -> bytes:
    title_xf = st.xf(font=st.font(True, "FF1F3864"))
    head_xf = st.xf(font=st.font(True, "FFFFFFFF"), fill=st.fill("FF1F3864"), wrap=True)
    kinds = {
        DATE: st.xf(num=st.numfmt(DATE_FMT)),
        NUM: st.xf(num=st.numfmt(NUM_FMT)),
        QTY: st.xf(num=st.numfmt(QTY_FMT)),
        INT: 0,
        TEXT: 0,
    }
    ncols = len(spec.columns)
    last_col = get_column_letter(ncols)
    last_row = spec.header_row + max(1, len(rows))
    W = etree.Element(q("worksheet"), nsmap={None: NS, "r": RNS})
    etree.SubElement(W, q("dimension")).set("ref", f"A1:{last_col}{last_row}")
    views = etree.SubElement(W, q("sheetViews"))
    v = etree.SubElement(views, q("sheetView"))
    v.set("workbookViewId", "0")
    pane = etree.SubElement(v, q("pane"))
    pane.set("ySplit", str(spec.header_row))
    pane.set("topLeftCell", f"A{spec.header_row + 1}")
    pane.set("activePane", "bottomLeft")
    pane.set("state", "frozen")
    etree.SubElement(W, q("sheetFormatPr")).set("defaultRowHeight", "14.4")
    cols = etree.SubElement(W, q("cols"))
    for i, c in enumerate(spec.columns, start=1):
        ce = etree.SubElement(cols, q("col"))
        ce.set("min", str(i))
        ce.set("max", str(i))
        ce.set("width", str(c.width))
        ce.set("customWidth", "1")
    sd = etree.SubElement(W, q("sheetData"))

    def text_cell(row_el, ref, text, style):
        c = etree.SubElement(row_el, q("c"))
        c.set("r", ref)
        c.set("t", "inlineStr")
        if style:
            c.set("s", str(style))
        is_ = etree.SubElement(c, q("is"))
        t = etree.SubElement(is_, q("t"))
        t.text = _escape_sheet_xml_text(text)
        if text != text.strip():
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

    r1 = etree.SubElement(sd, q("row"))
    r1.set("r", "1")
    text_cell(r1, "A1", spec.title, title_xf)
    r2 = etree.SubElement(sd, q("row"))
    r2.set("r", str(spec.header_row))
    r2.set("ht", "30")
    r2.set("customHeight", "1")
    for i, c in enumerate(spec.columns, start=1):
        text_cell(r2, f"{get_column_letter(i)}{spec.header_row}", c.header, head_xf)
    for n, rec in enumerate(rows, start=spec.header_row + 1):
        re_ = etree.SubElement(sd, q("row"))
        re_.set("r", str(n))
        for i, col in enumerate(spec.columns, start=1):
            v = rec.get(col.key)
            if v is None or v == "":
                continue
            ref = f"{get_column_letter(i)}{n}"
            if col.kind == TEXT or isinstance(v, str):
                text_cell(re_, ref, str(v), 0)
                continue
            c = etree.SubElement(re_, q("c"))
            c.set("r", ref)
            if kinds[col.kind]:
                c.set("s", str(kinds[col.kind]))
            if col.kind == DATE:
                val = to_serial(v)
            else:
                val = float(v)
            etree.SubElement(c, q("v")).text = str(int(val)) if float(val).is_integer() else repr(round(float(val), 6))
    af = etree.SubElement(W, q("autoFilter"))
    af.set("ref", f"A{spec.header_row}:{last_col}{last_row}")
    pm = etree.SubElement(W, q("pageMargins"))
    for k, val in (("left", "0.7"), ("right", "0.7"), ("top", "0.75"), ("bottom", "0.75"), ("header", "0.3"), ("footer", "0.3")):
        pm.set(k, val)
    return etree.tostring(W, xml_declaration=True, encoding="UTF-8", standalone=True)


def _ensure_sheet_part(wb: Workbook, name: str) -> str:
    """Return the part path of app sheet ``name``, registering a new one if needed."""
    if name in wb.sheet_paths:
        return wb.sheet_paths[name]
    existing = set(wb._parts)
    i = 1
    while f"xl/worksheets/sheet{i}.xml" in existing:
        i += 1
    path = f"xl/worksheets/sheet{i}.xml"
    # workbook.xml: <sheet> appended at the end (keeps existing localSheetId indexes)
    wbroot = wb._tree("xl/workbook.xml").getroot()
    sheets = wbroot.find("m:sheets", N)
    ids = [int(s.get("sheetId")) for s in sheets]
    rels_path = "xl/_rels/workbook.xml.rels"
    rels = wb._tree(rels_path).getroot()
    rid_nums = [int(re.sub(r"\D", "", r.get("Id")) or 0) for r in rels]
    rid = f"rId{max(rid_nums + [0]) + 1}"
    rel = etree.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship")
    rel.set("Id", rid)
    rel.set("Type", WS_REL)
    rel.set("Target", f"worksheets/sheet{i}.xml")
    s = etree.SubElement(sheets, q("sheet"))
    s.set("name", name)
    s.set("sheetId", str(max(ids + [0]) + 1))
    s.set(f"{{{RNS}}}id", rid)
    ct = wb._tree("[Content_Types].xml").getroot()
    o = etree.SubElement(ct, f"{{{CT_NS}}}Override")
    o.set("PartName", "/" + path)
    o.set("ContentType", WS_CT)
    for part in ("xl/workbook.xml", rels_path, "[Content_Types].xml"):
        wb._touch(part)
    wb.add_part(path, b"")
    wb.sheet_paths[name] = path
    return path


def write_tables(data: bytes, tables: Iterable[tuple[TableSpec, list[dict]]], *, protect: str) -> bytes:
    """Create/replace app sheets. ``protect`` is the name of the original tracking sheet:
    its XML, the shared strings and the calc chain must come out byte-identical."""
    wb = Workbook(data)
    if protect not in wb.sheet_paths:
        raise AppSheetError(f"Workbook has no sheet '{protect}'")
    protected_parts = {wb.sheet_paths[protect], "xl/sharedStrings.xml", "xl/calcChain.xml"}
    app_paths: set[str] = set()
    styles_before = etree.tostring(wb._tree("xl/styles.xml").getroot())
    st = _Styles(wb)
    for spec, rows in tables:
        if spec.sheet == protect:
            raise AppSheetError("Refusing to overwrite the original tracking sheet")
        path = _ensure_sheet_part(wb, spec.sheet)
        app_paths.add(path)
        wb._trees.pop(path, None)
        wb._parts[path] = _sheet_xml(spec, rows, st)
    wb._touch("xl/styles.xml")
    out = wb.to_bytes()
    _guard(data, out, protected_parts, app_paths, styles_before)
    return out


def _guard(old: bytes, new: bytes, protected: set[str], app_paths: set[str], styles_before: bytes) -> None:
    """Independent check: the original sheet and everything it depends on is untouched."""
    import io
    import zipfile

    zo, zn = zipfile.ZipFile(io.BytesIO(old)), zipfile.ZipFile(io.BytesIO(new))
    old_names = set(zo.namelist())
    for p in protected:
        if p in old_names and zo.read(p) != zn.read(p):
            raise AppSheetError(f"Safety check failed: {p} would change; nothing was written")
    allowed = {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml", "xl/styles.xml"}
    for p in old_names:
        if p in allowed or p in protected or p in app_paths:
            continue
        if zo.read(p) != zn.read(p):
            raise AppSheetError(f"Safety check failed: {p} would change; nothing was written")
    # styles.xml: only appended entries; every pre-existing child unchanged and in place
    old_root = etree.fromstring(styles_before)
    new_root = etree.fromstring(zn.read("xl/styles.xml"))
    for tag in ("numFmts", "fonts", "fills", "borders", "cellStyleXfs", "cellXfs", "cellStyles"):
        a, b = old_root.find(f"m:{tag}", N), new_root.find(f"m:{tag}", N)
        if a is None:
            continue
        if b is None or len(b) < len(a) or any(etree.tostring(x) != etree.tostring(y) for x, y in zip(a, b)):
            raise AppSheetError(f"Safety check failed: existing styles ({tag}) would change; nothing was written")
